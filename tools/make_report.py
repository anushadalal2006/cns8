"""Build the lab report (report/*.docx) from the measured data, in the order of Section 19.

usage: python3 tools/make_report.py      (then: soffice --headless --convert-to pdf ...)
"""
import csv
import os
import re
import statistics

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAME, UID = "Anusha Dalal", "2024300032"
ALGS = ["ecdsa", "rsapss", "mldsa", "slhdsa"]
LABEL = {"ecdsa": "ECDSA", "rsapss": "RSA-PSS", "mldsa": "ML-DSA", "slhdsa": "SLH-DSA"}
OUT = f"report/CNS_Exp08_Digital_Signatures_{UID}_{NAME.replace(' ', '_')}.docx"


# ------------------------------------------------------------------ data
def read_csv(path):
    with open(path) as f:
        return list(csv.DictReader(f))


def timing(rows):
    return {(r["algorithm"], r["operation"]): (float(r["mean_ms"]), float(r["std_ms"])) for r in rows}


SENDER = read_csv("sender/results.csv")
RECEIVER = read_csv("receiver/results.csv")
RUNS = {r: read_csv(f"sender/results_run{r}.csv") for r in (1, 2, 3)}
CPUS1 = read_csv("sender/results_cpus1.csv")
T = timing(SENDER)

sizes = {}
for line in open("logs/G_sizes.log"):
    m = re.match(r"(\w+) public=(\d+) private=(\d+) sig=(\d+) doc=(\d+)", line)
    if m:
        a, pk, sk, sg, doc = m.groups()
        sizes[a] = dict(pk=int(pk), sk=int(sk), sig=int(sg), doc=int(doc))
for a in ALGS:
    sizes[a]["tx"] = sizes[a]["doc"] + sizes[a]["sig"] + sizes[a]["pk"]
DOC_TEXT = open("sender/outbox/document.txt").read()
TAMPERED = open("receiver/inbox/document.txt.tampered", encoding="utf-8", errors="replace").read()
ENV = [l.strip() for l in open("sender/environment.txt") if l.strip()]
http_log = open("logs/F_http_server_log.log").read()
receiver_ip = re.search(r"^(\d+\.\d+\.\d+\.\d+) - -", http_log, re.M).group(1)

# Table 3 from the transcript: every verify block has a "# Test n" comment and a RESULT line
results3 = {}
cur = None
for line in open("logs/H_tamper.log"):
    m = re.search(r"verify\.py (\w+) .*# Test (4b|\d)", line)
    if m:
        cur = (m.group(2), m.group(1))
    elif "RESULT" in line and cur:
        results3[cur] = "PASS" if "PASS" in line else "FAIL"
        cur = None


def r(x, nd=2):
    return f"{x:.{nd}f}"


def ms(a, op):
    return T[(a, op)][0]


# ------------------------------------------------------------------ docx helpers
doc = Document()
st = doc.styles["Normal"]
st.font.name = "Calibri"
st.font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def H(text, level=1):
    doc.add_heading(text, level=level)


def P(text="", bold=False, italic=False, size=None, align=None):
    p = doc.add_paragraph()
    # **bold** fragments inside text
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        run = p.add_run(part)
        run.bold = bold or (i % 2 == 1)
        run.italic = italic
        if size:
            run.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def B(text):
    p = P(text)
    p.style = doc.styles["List Bullet"]
    return p


def table(header, rows, widths=None, size=9.5, first_col_bold=True):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h)); run.bold = True; run.font.size = Pt(size)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(size)
            if i == 0 and first_col_bold:
                run.bold = True
            if str(v).startswith("PASS"):
                run.font.color.rgb = RGBColor(0x1E, 0x7B, 0x34)
            elif str(v).startswith("FAIL"):
                run.font.color.rgb = RGBColor(0xB0, 0x2A, 0x2A)
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
            t._tbl.tblGrid.gridCol_lst[i].w = Cm(w)
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph()
    return t


def caption(text):
    P(text, italic=True, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)


def image(path, width=16.5, cap=None):
    doc.add_picture(path, width=Cm(width))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if cap:
        caption(cap)


def code(text, size=7.5):
    for chunk in text.rstrip("\n").split("\n"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        run = p.add_run(chunk if chunk else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(size)
    doc.add_paragraph()


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ------------------------------------------------------------------ title page
P("Bharatiya Vidya Bhavan's", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
P("SARDAR PATEL INSTITUTE OF TECHNOLOGY", bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Department of Computer Engineering - Cryptography and Network Security Lab",
  align=WD_ALIGN_PARAGRAPH.CENTER)
P("Experiment 8", bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
P("Evaluation of Classical and Post-Quantum Digital Signature Algorithms using Docker-Based "
  "Client-Server Communication", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
table(["Field", "Value"], [["Name", NAME], ["UID", UID], ["Date", "16 October 2026"],
                           ["Algorithms", "RSA-PSS (RSA-3072), ECDSA (P-256), ML-DSA-44, SLH-DSA-SHA2-128s"]],
      widths=[4, 12])
H("Aim", 2)
P("To implement the complete digital-signature workflow (key generation, signing, transmission, "
  "verification and tamper detection) between two Docker containers acting as sender and receiver, "
  "for two classical algorithms (RSA-PSS and ECDSA) and two NIST-standardised post-quantum algorithms "
  "(ML-DSA and SLH-DSA), and to compare experimentally their key sizes, signature sizes, speed and "
  "communication overhead, and to study why post-quantum signatures became necessary.")
H("UID-based individualisation (Section 5)", 2)
table(["Item", "Value for UID " + UID],
      [["Document line 1", DOC_TEXT.splitlines()[0]],
       ["Document line 2", DOC_TEXT.splitlines()[1]],
       ["A = 1000 + (UID mod 9000)", f"1000 + {int(UID) % 9000} = {1000 + int(UID) % 9000}"],
       ["ECDSA private scalar d", "d = (int(SHA-256(UID || \"CNS-DSIG-ECDSA-KEY\")) mod (n-1)) + 1 = "
        "0x3225474890762ddbc056efc9926a6ed23cf2e7b72731af6acd57e5949ab4f561"],
       ["Benchmark message", "random.Random(2024300032).randbytes(1024), SHA-256 prefix c0aa6017dac0b92a"],
       ["Tamper index (document, 59 B)", f"{UID} mod 59 = {int(UID) % 59}"],
       ["Tamper index (signatures)", "ECDSA: mod 71 = 4, RSA-PSS: mod 384 = 256, ML-DSA: mod 2420 = 1492, "
        "SLH-DSA: mod 7856 = 5232"]],
      widths=[5, 12])
page_break()

# ------------------------------------------------------------------ 1. Part A
H("1. Part A - Hand-drawn digital-signature diagram")
P("**[ATTACH THE HAND-DRAWN SHEET HERE]** - Part A must be drawn with pen/pencil and a ruler on a "
  "plain/ruled sheet before using the computer, with name, UID and date in the top-right corner, and "
  "appended to this document. The checklist below lists the eight required elements; a computer-drawn "
  "reference of the same diagram (to copy from) is given in Appendix A.")
table(["No.", "Required element", "Naming / convention used", "Done"],
      [["1", "Two parties and an insecure channel", "Sender (Alice) left, Receiver (Bob) right, channel between", "[ ]"],
       ["2", "Message and hash", "Message M; hash function H( ); digest h = H(M)", "[ ]"],
       ["3", "Sender's key pair", "Private key SK_A (secret, dashed arrow), public key PK_A", "[ ]"],
       ["4", "Signing step", "S = Sign(SK_A, h); output pair (M || S)", "[ ]"],
       ["5", "Public-key distribution", "Arrow PK_A to receiver; SK_A is never sent", "[ ]"],
       ["6", "Receiver side", "Received M' and S; h' = H(M'); Verify(PK_A, h', S)", "[ ]"],
       ["7", "Decision", "Decision box: ACCEPT (valid) / REJECT (invalid)", "[ ]"],
       ["8", "Attacker case", "Mallory modifies M to M' in the channel -> REJECT", "[ ]"]],
      widths=[1.2, 4.5, 9.5, 1.3])
page_break()

# ------------------------------------------------------------------ 2. Part B/C
H("2. Experiment architecture and Docker environment (Parts B and C)")
P("One image (sig-lab:1.0) and two containers, sender (Alice) and receiver (Bob), attached to the "
  "user-defined bridge network sig-net. The host folders sender/ and receiver/ are bind-mounted at /lab. "
  "The sender generates keys, signs the document and serves outbox/ over HTTP on port 8000; the receiver "
  "downloads document, signature and public key with curl and verifies them. Only public information "
  "crosses the network; private keys stay in the sender's keys/ folder.")
P(f"IP addresses (Step C7): sender = 172.18.0.2, receiver = 172.18.0.3 (subnet 172.18.0.0/16 of sig-net).")
H("Table 1: Execution environment", 2)
table(["CPU model", "Cores / threads", "RAM", "Host OS", "Docker version", "cryptography version",
       "liboqs version"],
      [["Intel(R) Xeon(R) Processor @ 2.80GHz", "4 vCPUs (4 threads)", "15 GiB",
        "Ubuntu 24.04.5 LTS, kernel 6.18.44", "29.8.2 (build 7fc2dff)", "50.0.2",
        "liboqs 0.16.0 (liboqs-python 0.16.0.1)"]],
      size=8.5, first_col_bold=False)
P("Container: Ubuntu 24.04.5 LTS, Python 3.12.3. The machine is a 4-vCPU cloud virtual machine. "
  "Note: the build machine reaches the Internet through a TLS-inspecting proxy, so the image was "
  "built from a copy of the Dockerfile that additionally trusts that proxy's CA certificate "
  "(Dockerfile.sandbox in the repository); the submitted Dockerfile is exactly the one in Step C3.",
  size=9, italic=True)
image("screenshots/01_docker_network_create.png", 16.5, "Screenshot 1: Step C5 - docker network create / ls / inspect sig-net")
image("screenshots/02_docker_ps.png", 16.5, "Screenshot 2: Step C6 - both containers started and Up (docker ps)")
image("screenshots/03_network_inspect_sig_net.png", 13, "Screenshot 3: Step C7 - docker network inspect sig-net (both containers attached)")
image("screenshots/04_ping_test.png", 16.5, "Screenshot 4: Step C7 - ping test sender -> receiver, STUDENT_UID, liboqs mechanisms, IP addresses")
image("screenshots/05_environment.png", 16.5, "Screenshot 5: Step C8 - execution environment")
page_break()

# ------------------------------------------------------------------ 3. screenshots D-G
H("3. Key generation, signing, transfer and verification (Parts D-G)")
image("screenshots/06_alg_names.png", 16.5, "Screenshot 6: Step D1 - exact algorithm names: ML-DSA-44 and SLH_DSA_PURE_SHA2_128S")
image("screenshots/07_keygen.png", 15, "Screenshot 7: Step D3 - keygen.py for all four algorithms")
image("screenshots/08_signing.png", 15, "Screenshot 8: Part E - sign.py for all four algorithms, document and its SHA-256")
image("screenshots/09_transfer_receiver.png", 16.5, "Screenshot 9: Part F - receiver downloads over sig-net; SHA-256 equals the sender's value")
image("screenshots/10_transfer_sender_http_log.png", 15,
      f"Screenshot 10: Part F - sender's HTTP server log; every request comes from the receiver's IP {receiver_ip}")
image("screenshots/11_tcpdump.png", 16.5, "Screenshot 11: Part F (optional) - tcpdump on the receiver, port 8000")
image("screenshots/12_verification.png", 13, "Screenshot 12: Part G - verify.py prints PASS for all four algorithms")
P("Document SHA-256 on the sender and the receiver: d44e70dfc1a5a1f16b40f91fc9c652fc2a9e60d84dbf74435a3f4deb7cf7f388 "
  "(identical, so the file was transferred unchanged).")
page_break()

# ------------------------------------------------------------------ 4. Tables 2, 3, 4
H("4. Results: sizes, tampering and timings")
H("Table 2: Key, signature and transmission sizes (bytes)", 2)
table(["Algorithm", "Public key", "Private key", "Signature", "Document", "Transmitted (doc + sig + public key)"],
      [[LABEL[a], sizes[a]["pk"], sizes[a]["sk"], sizes[a]["sig"], sizes[a]["doc"], sizes[a]["tx"]] for a in ALGS])
P("Measured with stat -c %s (Screenshot 13). Sanity check: ML-DSA-44 public key 1312, secret key 2560, "
  "signature 2420 bytes and SLH-DSA-SHA2-128s public key 32, secret key 64, signature 7856 bytes - exactly "
  "the FIPS 204/205 values, since liboqs stores raw keys. The classical keys are DER encoded and therefore "
  "slightly larger than the raw values (P-256: 65-byte point + 26-byte SubjectPublicKeyInfo header = 91 B; "
  "ECDSA signatures are DER and vary between 70 and 72 B, 71 B here; an RSA-3072 PKCS#8 key varies by "
  "about 1 B between key pairs, e.g. 1792 B here and 1793 B in the benchmark).", size=9.5)
image("screenshots/13_sizes.png", 16.5, "Screenshot 13: file sizes on the sender and bytes received by the receiver")

H("Table 3: Verification results", 2)
rows3 = []
names = {"1": "1. Original document + original signature", "2": "2. One character of document changed",
         "3": "3. One bit of signature changed", "4": "4. Correct document + signature, wrong public key (receiver's own key pair)"}
for t, nm in names.items():
    rows3.append([nm] + [results3.get((t, a), "-") for a in ALGS] + ["PASS" if t == "1" else "FAIL"])
rows3.append(["4b. (supplementary) ECDSA public key derived from a different UID (2024300033)",
              results3.get(("4b", "ecdsa"), "-"), "n/a", "n/a", "n/a", "FAIL"])
table(["Test case", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA", "Expected"], rows3,
      widths=[6.5, 2, 2, 2, 2, 2])
P("**Original and modified document (Test 2).** index = UID mod 59 = 2, byte 0x61 ('a') -> 0x60 ('`'):")
table(["", "Document text"],
      [["Original M", DOC_TEXT.replace("\n", " / ").rstrip(" /")],
       ["Modified M'", TAMPERED.replace("\n", " / ").rstrip(" /")]], widths=[3, 13])
P("**Signature bit flipped (Test 3)** at index UID mod length: ECDSA byte 4 (0x00 -> 0x01), RSA-PSS "
  "byte 256 (0x9b -> 0x9a), ML-DSA byte 1492 (0x09 -> 0x08), SLH-DSA byte 5232 (0x53 -> 0x52).")
P("**Observation on Test 4 (ECDSA = PASS).** In Test 4 the receiver runs keygen.py ecdsa. The ECDSA private "
  "key is not random: as required by Section 5 it is derived from the UID, and both containers receive "
  "the same STUDENT_UID. The receiver therefore regenerates exactly the same key pair as the sender (cmp "
  "confirms the two ecdsa_public.key files are byte-identical, Screenshot 16), so the \"wrong\" key is in fact "
  "the right one and verification correctly passes. This is precisely the weakness of the warning box in "
  "Section 5: anyone who knows (or brute-forces, ~33 bits) the UID can recreate the private key and forge "
  "signatures. RSA-PSS, ML-DSA and SLH-DSA keys come from the secure random generator, so the receiver's "
  "key pairs differ and verification fails as expected. With a genuinely different ECDSA key (Test 4b, key "
  "derived from UID 2024300033) ECDSA also returns FAIL.")
image("screenshots/14_tamper_test1_test2.png", 15, "Screenshot 14: Part H - tamper.py doc, Tests 1 and 2")
image("screenshots/15_tamper_test3.png", 15, "Screenshot 15: Part H - Test 3 (one bit of each signature flipped)")
image("screenshots/16_tamper_test4.png", 15, "Screenshot 16: Part H - Test 4 (receiver's own key pair) and supplementary Test 4b")

H("Table 4: Timing results (milliseconds, mean ± standard deviation)", 2)
P("Sender container, median of three bench.py runs (run 2; runs ranked by total mean time). Each value is "
  "the mean ± standard deviation over the individual calls: key generation n = 20, signing n = 100, "
  "verification n = 100, after 5 warm-up calls; 1024-byte UID-seeded message.", size=9.5)
table(["Algorithm", "Key generation", "Signing", "Verification", "Sign / Verify ratio"],
      [[LABEL[a]] + [f"{T[(a, op)][0]:.4f} ± {T[(a, op)][1]:.4f}" for op in ("keygen", "sign", "verify")]
       + [f"{ms(a, 'sign') / ms(a, 'verify'):.2f}"] for a in ALGS])
P("Supplementary - variation between the three runs on the sender (mean of the three run means ± their "
  "standard deviation), the receiver container, and the sender limited to one CPU (docker update --cpus 1):",
  size=9.5)
sup = []
for a in ALGS:
    row = [LABEL[a]]
    for op in ("keygen", "sign", "verify"):
        vals = [timing(RUNS[k])[(a, op)][0] for k in (1, 2, 3)]
        row.append(f"{statistics.mean(vals):.4f} ± {statistics.stdev(vals):.4f}")
    row += [f"{timing(RECEIVER)[(a, op)][0]:.4f}" for op in ("keygen", "sign", "verify")]
    row += [f"{timing(CPUS1)[(a, 'sign')][0]:.4f}"]
    sup.append(row)
table(["Alg.", "KeyGen (3 runs)", "Sign (3 runs)", "Verify (3 runs)", "KeyGen recv.", "Sign recv.",
       "Verify recv.", "Sign --cpus 1"], sup, size=8)
P("The sender and receiver containers give the same picture (same image, same host): differences are "
  "within run-to-run noise. Limiting the container to one CPU changes little because each operation is "
  "single-threaded. RSA key generation shows a very large standard deviation because the time to find "
  "two random 1536-bit primes is itself random.", size=9.5)
image("screenshots/17_bench_sender.png", 15, "Screenshot 17: Part I - bench.py on the sender (three runs)")
image("screenshots/18_bench_receiver.png", 15, "Screenshot 18: Part I - bench.py on the receiver")
image("screenshots/19_bench_sender_cpus1.png", 15, "Screenshot 19: Part I (optional) - sender limited to 1 CPU, median-run selection")
page_break()

# ------------------------------------------------------------------ 5. Graphs and Tables 5-7
H("5. Graphical analysis (Section 15)")
env_caption = "Environment: " + "; ".join(ENV) + f". UID {UID}."
graphs = [
    ("graph1_key_sizes.png", "Graph 1: Public-key and private-key size (bytes, log scale)"),
    ("graph2_signature_size.png", "Graph 2: Signature size (bytes, log scale)"),
    ("graph3_signing_time.png", "Graph 3: Average signing time (ms, log scale), error bar = standard deviation, n = 100"),
    ("graph4_verification_time.png", "Graph 4: Average verification time (ms, log scale), error bar = standard deviation, n = 100"),
    ("graph5_keygen_time.png", "Graph 5: Average key-generation time (ms, log scale), n = 20"),
    ("graph6_transmitted_bytes.png", "Graph 6: Total transmitted bytes = document + signature + public key (stacked)"),
    ("graph7_relative_to_ecdsa.png", "Graph 7 (extra): Key-generation, signing and verification time relative to ECDSA (ECDSA = 1, log scale)"),
]
for g, cap in graphs:
    image(os.path.join("graphs", g), 14.5, f"{cap}. {env_caption}")

H("6. Comparative analysis (Part J)")
H("Table 5: Relative comparison with ECDSA as the baseline", 2)
def rel(vals):
    return ["1.0"] + [f"{vals[a] / vals['ecdsa']:.2f}" for a in ALGS[1:]]
table(["Parameter", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA"],
      [["Public key size"] + rel({a: sizes[a]["pk"] for a in ALGS}),
       ["Signature size"] + rel({a: sizes[a]["sig"] for a in ALGS}),
       ["Signing time"] + rel({a: ms(a, "sign") for a in ALGS}),
       ["Verification time"] + rel({a: ms(a, "verify") for a in ALGS}),
       ["Key generation time"] + rel({a: ms(a, "keygen") for a in ALGS}),
       ["Total transmitted bytes"] + rel({a: sizes[a]["tx"] for a in ALGS})])

H("Table 6: Property comparison", 2)
table(["Property", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA"],
      [["Standard", "FIPS 186-5", "RFC 8017, FIPS 186-5", "FIPS 204 (2024)", "FIPS 205 (2024)"],
       ["Security parameter / category", "P-256, ~128-bit classical", "RSA-3072, ~128-bit classical",
        "ML-DSA-44, NIST category 2", "SLH-DSA-SHA2-128s, NIST category 1"],
       ["Mathematical basis", "Elliptic-curve discrete logarithm", "Integer factorisation",
        "Module lattices (Module-LWE / Module-SIS)", "Hash functions only (hypertree of XMSS/WOTS+ and FORS trees)"],
       ["Resistant to Shor's algorithm?", "No", "No", "Yes", "Yes"],
       ["Public key size", f"Small ({sizes['ecdsa']['pk']} B)", f"Medium ({sizes['rsapss']['pk']} B)",
        f"Large ({sizes['mldsa']['pk']} B)", f"Very small ({sizes['slhdsa']['pk']} B)"],
       ["Signature size", f"Small ({sizes['ecdsa']['sig']} B)", f"Medium ({sizes['rsapss']['sig']} B)",
        f"Large ({sizes['mldsa']['sig']} B)", f"Very large ({sizes['slhdsa']['sig']} B)"],
       ["Fastest operation", f"Signing ({r(ms('ecdsa','sign'),3)} ms)", f"Verification ({r(ms('rsapss','verify'),3)} ms)",
        f"Verification ({r(ms('mldsa','verify'),3)} ms)", f"Verification ({r(ms('slhdsa','verify'),3)} ms)"],
       ["Slowest operation", f"Verification ({r(ms('ecdsa','verify'),3)} ms)", f"Key generation ({r(ms('rsapss','keygen'),1)} ms)",
        f"Signing ({r(ms('mldsa','sign'),3)} ms)", f"Signing ({r(ms('slhdsa','sign'),1)} ms)"],
       ["Maturity / deployment", "Very mature; TLS, SSH, code signing, smart cards, cryptocurrencies",
        "Very mature (decades); PKI, TLS, document signing", "New (2024); being rolled out in libraries, HSMs, hybrid TLS/PKI",
        "New standard (2024) but very conservative design (SPHINCS+); firmware, root keys"],
       ["Main advantage", "Smallest keys and signatures, fast signing", "Very fast verification, universally supported",
        "Quantum-safe with the best overall speed; fastest verification measured", "Quantum-safe with security resting only on hash functions; 32-byte public key"],
       ["Main disadvantage", "Broken by Shor; needs good randomness/nonces", "Broken by Shor; slow key generation, large keys for its security",
        "Large public key and signature (about 3.8 KB per message)", "Very slow signing and very large signatures"]],
      size=8.5, widths=[3.2, 3.3, 3.3, 3.4, 3.4])

H("Table 7: Which algorithm would you choose?", 2)
table(["Scenario", "Chosen algorithm", "Justification"],
      [["Signing a firmware image for a device that stays in the field for 15 years", "SLH-DSA",
        f"Must survive a future quantum computer; hash-only security is the most conservative, signing is done "
        f"rarely at the vendor (840 ms is fine) and device verification is cheap ({r(ms('slhdsa','verify'),2)} ms) with a 32 B key to store."],
       ["Authenticating a short sensor message on a very low-bandwidth link", "ECDSA",
        f"Smallest overhead ({sizes['ecdsa']['sig']} B signature, {sizes['ecdsa']['pk']} B key); a short-lived sensor "
        f"reading needs no decades-long quantum resistance, while ML-DSA would add {sizes['mldsa']['sig']} B per message."],
       ["Signing a TLS server certificate in 2026 on a general web server", "ECDSA (P-256), moving to hybrid ECDSA + ML-DSA",
        "Universally supported by browsers today with tiny handshakes; certificates live < 1 year, so authentication "
        "is not exposed to harvest-now-forge-later, and ML-DSA can be added as clients support it."],
       ["Root key of a certificate authority that must remain trusted for 30 years", "SLH-DSA",
        "Security relies only on hash functions (no new structured assumption); a root signs rarely so slow signing and "
        "7.8 KB signatures are acceptable, and it must remain unforgeable well into the quantum era."],
       ["High-volume server that verifies millions of signatures per second", "ML-DSA",
        f"Fastest verification measured ({r(ms('mldsa','verify'),3)} ms vs ECDSA {r(ms('ecdsa','verify'),3)} ms and "
        f"RSA {r(ms('rsapss','verify'),3)} ms) and quantum-safe; RSA-PSS verifies almost as fast but is not post-quantum."]],
      size=9, widths=[5, 3, 9])
page_break()

# ------------------------------------------------------------------ 7. Discussion
H("7. Discussion questions")
sig_ratio = sizes["slhdsa"]["sig"] / sizes["ecdsa"]["sig"]
QA = [
 ("Conceptual", None),
 ("1. What are the three algorithms of a digital-signature scheme? Which key is used in which algorithm?",
  "KeyGen() -> (PK, SK) creates the key pair and uses no key as input (only randomness; in this lab the ECDSA "
  "scalar is derived from the UID). Sign(SK, M) -> S uses the private key SK, known only to the signer. "
  "Verify(PK, M, S) -> accept/reject uses the public key PK, which anyone may hold. In hash-and-sign schemes "
  "(ECDSA, RSA-PSS) Sign and Verify work on h = H(M); ML-DSA and SLH-DSA take the whole message and hash it "
  "internally, which is why sign.py passes the document directly to them."),
 ("2. What is the difference between a digital signature and a MAC (for example HMAC)? Why does the signature allow non-repudiation?",
  "A MAC uses one shared secret key for both creating and checking the tag, so it gives integrity and "
  "authentication only between parties that share the key; since the verifier can compute the same tag, a third "
  "party cannot tell who produced it. A signature uses an asymmetric key pair: only the holder of SK can "
  "produce S, while anyone with PK can verify. Because the verifier cannot create signatures, a valid signature "
  "is evidence to a judge or third party that the SK holder signed the message - the signer cannot later deny it "
  "(non-repudiation). MACs are, however, much faster and shorter."),
 ("3. Why is the private key never transferred in this experiment? What would happen if it were?",
  "Verification needs only PK, and the security of the whole scheme rests on SK being known to the signer alone. "
  "Our channel (plain HTTP over sig-net) is readable by anyone on the network - the tcpdump capture shows the "
  "file contents in clear. If SK were sent, any eavesdropper (Mallory) could sign arbitrary documents such as "
  "\"Transfer Rs.999999\" that verify as PASS; integrity, authentication and non-repudiation would all be lost "
  "and the key pair would have to be revoked. This is why outbox/ contains only public files and keys/ is never "
  "served."),
 ("4. Why did post-quantum cryptography come into the picture? Explain Shor's algorithm's effect on RSA and ECDSA, and why a signature that is valid today may be dangerous to rely on in 20 years.",
  "RSA is secure because factoring n = p·q is hard, and ECDSA because computing d from Q = d·G (the elliptic-curve "
  "discrete logarithm) is hard. Shor's algorithm on a large fault-tolerant quantum computer solves both problems "
  "in polynomial time, so an attacker could compute the private key directly from the public key and forge any "
  "signature; increasing key sizes does not help. Grover's algorithm only gives a quadratic speed-up against hash "
  "functions, and no efficient quantum attack is known on lattice problems - hence ML-DSA (lattices) and SLH-DSA "
  "(hashes), standardised by NIST in August 2024. Signatures that must stay trustworthy for decades (firmware, "
  "root certificates, legal archives, land records) are the concern: once a quantum computer exists, an attacker "
  "who recovers the old RSA/ECDSA key can create new signatures that are indistinguishable from the genuine "
  "ones, so a verifier in 20 years can no longer tell a real 2026 contract from a forged one. Long-lived systems "
  "and their keys must therefore migrate before such a machine exists."),
 ("Analytical (measured data)", None),
 ("5. Which algorithm had the smallest and which the largest signature? Express the difference as a ratio and explain the cause.",
  f"Smallest: ECDSA, {sizes['ecdsa']['sig']} B (DER encoding of two 32-byte integers r, s). Largest: SLH-DSA-SHA2-128s, "
  f"{sizes['slhdsa']['sig']} B - a ratio of {sig_ratio:.1f}x (and {sizes['slhdsa']['sig']/sizes['mldsa']['sig']:.2f}x "
  f"ML-DSA's {sizes['mldsa']['sig']} B, {sizes['slhdsa']['sig']/sizes['rsapss']['sig']:.1f}x RSA's {sizes['rsapss']['sig']} B). "
  "ECDSA is compact because elliptic-curve groups give 128-bit security with 256-bit numbers. An SLH-DSA signature "
  "must contain a FORS few-time signature plus a WOTS+ one-time signature and a Merkle authentication path for every "
  "one of the hypertree layers - many hash values of 16 bytes each. ML-DSA's signature is a vector of polynomials "
  "with 256 coefficients each (plus a hint), which is why it is in between."),
 ("6. Which algorithm was fastest to sign and which to verify? Why is SLH-DSA signing much slower than ML-DSA signing?",
  f"Fastest signing: ECDSA ({r(ms('ecdsa','sign'),4)} ms), closely followed by ML-DSA ({r(ms('mldsa','sign'),4)} ms). "
  f"Fastest verification: ML-DSA ({r(ms('mldsa','verify'),4)} ms), followed by RSA-PSS ({r(ms('rsapss','verify'),4)} ms, "
  f"small exponent e = 65537) and ECDSA ({r(ms('ecdsa','verify'),4)} ms). SLH-DSA signing took {r(ms('slhdsa','sign'),1)} ms, "
  f"about {ms('slhdsa','sign')/ms('mldsa','sign'):,.0f} times ML-DSA. ML-DSA signing is a few polynomial multiplications "
  "(NTT) with occasional rejection-sampling retries. SLH-DSA is stateless, so each signature must regenerate, from the "
  "secret seed, the parts of a huge hypertree of hash-based one-time keys: building the FORS trees and the WOTS+ "
  "chains and Merkle trees of every layer costs hundreds of thousands of SHA-256 calls. The 's' (small) parameter set "
  "makes this worse because it uses tall trees to shrink the signature; verification only recomputes one path, so it "
  f"is fast ({r(ms('slhdsa','verify'),3)} ms)."),
 ("7. SLH-DSA has a tiny public key (32 bytes) but a large signature. ML-DSA has a larger public key but a smaller signature. Discuss the trade-off for a system that sends one public key but millions of signatures.",
  f"If the public key is sent once and signatures are sent N times, the cost is PK + N·|S|. For ML-DSA it is "
  f"1312 + 2420·N bytes and for SLH-DSA 32 + 7856·N bytes. SLH-DSA is cheaper only for N = 0; already with one "
  f"signature ML-DSA's total (3732 B) is smaller than SLH-DSA's (7888 B), and for 1 million signatures ML-DSA needs "
  f"about 2.42 GB against 7.86 GB for SLH-DSA (3.2x more bandwidth/storage) - the public-key size becomes irrelevant. "
  "Add to that SLH-DSA's ~8000x slower signing. So for high-volume signing ML-DSA is clearly preferable; SLH-DSA's "
  "tiny key matters where the key must be stored in very constrained places (ROM, fuse bits of a boot loader, a "
  "certificate chain) and few signatures are produced, e.g. firmware and root keys."),
 ("8. The public key was fetched over the same untrusted channel as the document. Explain how a man-in-the-middle could defeat the scheme and how certificates and a PKI solve this.",
  "Mallory, sitting between sender and receiver (or impersonating the host name 'sender'), can replace all three files: "
  "she generates her own key pair, writes her own document, signs it with her SK and delivers her PK. The receiver "
  "verifies Mallory's signature with Mallory's key and gets PASS - verification only proves that the document matches "
  "the public key, not who owns the key. The fix is to authenticate the public key itself: a certificate authority "
  "signs a certificate binding PK_A to Alice's identity, the receiver checks the chain up to a root CA key that it "
  "obtained out of band (pre-installed trust store) and checks validity and revocation. Other ways are key pinning or "
  "comparing a key fingerprint over a separate trusted channel. In practice the transport is also protected by TLS, "
  "itself authenticated with certificates."),
 ("Application and evaluation", None),
 ("9. Is ML-DSA or SLH-DSA \"better\" than ECDSA? Answer using your data, and consider size, speed, maturity and quantum security. Is there a single winner?",
  f"There is no single winner. Size: ECDSA wins clearly - {sizes['ecdsa']['tx']} B transmitted vs "
  f"{sizes['mldsa']['tx']} B for ML-DSA ({sizes['mldsa']['tx']/sizes['ecdsa']['tx']:.1f}x) and {sizes['slhdsa']['tx']} B for "
  f"SLH-DSA ({sizes['slhdsa']['tx']/sizes['ecdsa']['tx']:.1f}x). Speed: ML-DSA is as fast as or faster than ECDSA "
  f"(key generation {ms('mldsa','keygen')/ms('ecdsa','keygen'):.2f}x, signing {ms('mldsa','sign')/ms('ecdsa','sign'):.2f}x, verification "
  f"{ms('mldsa','verify')/ms('ecdsa','verify'):.2f}x of ECDSA), while SLH-DSA signs ~{ms('slhdsa','sign')/ms('ecdsa','sign'):,.0f}x slower. "
  "Maturity: ECDSA has decades of analysis and universal support; the PQ standards are from 2024 and implementations "
  "are still being hardened. Quantum security: only ML-DSA and SLH-DSA survive Shor; SLH-DSA has the most conservative "
  "assumption. So ECDSA is best today where size matters and data is short-lived; ML-DSA is the best general-purpose "
  "PQ replacement; SLH-DSA is the choice for rarely-signed, long-lived roots of trust - and hybrids bridge the transition."),
 ("10. Why do many experts recommend hybrid signatures (a classical plus a post-quantum signature) during the transition period?",
  "A hybrid signature is valid only if both component signatures verify, so it stays secure as long as at least one "
  "of the two algorithms is unbroken. The PQ schemes are new: cryptanalysis or implementation bugs (side channels, "
  "incorrect rejection sampling) may still be found, and some PQ candidates (e.g. Rainbow, SIKE) were broken shortly "
  "before or during standardisation. The classical scheme is well understood but falls to a future quantum computer. "
  "Combining them protects against both risks, satisfies compliance rules that still require FIPS-186 algorithms, "
  "and lets systems interoperate with old verifiers during migration. The cost is the sum of both sizes and times "
  f"(e.g. ECDSA + ML-DSA: {sizes['ecdsa']['sig'] + sizes['mldsa']['sig']} B of signatures)."),
 ("11. Does the Docker network affect the cryptographic measurements? How did you keep Docker, the network and file I/O out of your timings, and what other factors can still influence them?",
  "No - the network and Docker only affect the transfer, not the sign/verify computation. bench.py generates the "
  "1024-byte message in memory, loads/creates keys before the timed region and wraps only the single library call "
  "(generate / sign / verify) between two time.perf_counter_ns() readings; there is no file access or network I/O "
  "in the loop, containers are already running, 5 warm-up calls are discarded, many repetitions give mean, std, min "
  "and max, and the last signature is verified so failures are not timed. Docker containers are ordinary processes "
  "using the host kernel, so CPU-bound code runs at native speed - the sender and receiver results agree within "
  "noise. Remaining influences: CPU frequency scaling / turbo and thermal state, other load on the host (this is a "
  "shared cloud VM - see the outliers in max_ms), scheduler pre-emption and cache effects, a container CPU limit "
  "(--cpus 1 changed little for these single-threaded operations), Python call overhead, library versions and "
  "build flags (AVX2 code paths in liboqs and OpenSSL), and the randomness of RSA prime generation."),
]
for q, a in QA:
    if a is None:
        H(q, 2)
    else:
        P(q, bold=True)
        P(a)
P("Related questions referred to in the lab sheet:", bold=True)
P("**Think box (Part F) - public key from the untrusted channel:** see Question 8 - in practice the public key is "
  "delivered inside a certificate signed by a trusted CA (PKI), or pinned / checked by fingerprint out of band.")
P("**UID-derived keys (warning in Section 5):** d is derived from a 10-digit UID, i.e. only about 10^10 ≈ 2^33 "
  "possibilities, so an attacker can try every UID in seconds and recover the ECDSA private key. Our Test 4 shows "
  "the consequence directly: the receiver re-created the sender's private key just from the UID. Real private keys "
  "must come from a cryptographically secure random generator (os.urandom / the OS CSPRNG) with the full 256-bit "
  "entropy, and be stored protected (HSM, smart card or encrypted PKCS#8).")
page_break()

# ------------------------------------------------------------------ 8. Conclusion
H("8. Final conclusion")
CONCLUSION = (
 "A digital signature gives integrity and authentication because only the holder of the private key can produce a "
 "signature, while anyone holding the matching public key can check it; the private key never left the sender's "
 "container, and only the document, signature and public key crossed the sig-net network. Changing a single "
 "character of the document or a single bit of the signature made verification fail for all four algorithms, as did "
 "verifying with an unrelated public key. The ECDSA case was instructive: because its key was derived from my UID, "
 "the receiver regenerated the same key and verification passed, showing why low-entropy keys are dangerous. "
 "ECDSA had the smallest footprint (71-byte signature, 221 bytes transmitted), RSA-PSS verified quickly but needed "
 "slow key generation, ML-DSA-44 was as fast as or faster than ECDSA but transmitted 3,791 bytes, and SLH-DSA had a "
 "32-byte public key but a 7,856-byte signature and about 840 ms signing. Post-quantum signatures are necessary "
 "because Shor's algorithm would break RSA and ECDSA, letting attackers forge signatures on documents that must stay "
 "trustworthy for decades; the price is larger keys and signatures and, for SLH-DSA, slower signing. For "
 "today's bandwidth-sensitive applications I would choose ECDSA; for general "
 "post-quantum use ML-DSA; and for long-term roots of trust and firmware SLH-DSA, whose security rests only on hash "
 "functions. I also learned to build one image, run isolated containers on a user-defined network, reach them by "
 "name, use bind mounts to keep results, and keep container overhead out of measurements.")
P(CONCLUSION)
P(f"(Word count: {len(CONCLUSION.split())})", italic=True, size=9)
page_break()

# ------------------------------------------------------------------ 9. Source code
H("9. Source code")
for f in ["Dockerfile", "sender/common.py", "sender/keygen.py", "sender/sign.py", "sender/verify.py",
          "sender/tamper.py", "sender/bench.py", "sender/plot.py"]:
    H(os.path.basename(f), 2)
    code(open(f).read())
P("The same programs are copied to the receiver folder (cp ~/sig-lab/sender/*.py ~/sig-lab/receiver/). The whole "
  "experiment can be repeated with run_experiment.sh, which also records every command transcript in logs/.",
  italic=True, size=9)
page_break()

# ------------------------------------------------------------------ 10. results.csv
H("10. results.csv files and UID")
P(f"UID: **{UID}** (written in the first column of every result file).")
cols = ["uid", "algorithm", "operation", "runs", "mean_ms", "std_ms", "min_ms", "max_ms",
        "public_key_bytes", "private_key_bytes", "signature_bytes"]
short = ["uid", "algorithm", "operation", "runs", "mean_ms", "std_ms", "min_ms", "max_ms", "pk_B", "sk_B", "sig_B"]
for title, rows in (("Sender - results.csv (median run = results_run2.csv)", SENDER),
                    ("Receiver - results.csv", RECEIVER)):
    H(title, 2)
    table(short, [[row[c] for c in cols] for row in rows], size=7.5, first_col_bold=False)
P("Header line of the CSV files: " + ",".join(cols) + ". Also submitted: results_run1.csv, results_run2.csv, "
  "results_run3.csv and results_cpus1.csv (sender).", size=9)

# ------------------------------------------------------------------ Appendix A
page_break()
H("Appendix A - Reference digital-signature diagram (for the hand-drawn Part A)")
P("Computer-drawn reference containing all eight required elements. It does NOT replace the hand-drawn "
  "sheet required by Part A.", italic=True, size=9)
image("report/partA_reference_diagram.png", 17)

doc.save(OUT)
print("wrote", OUT, "| conclusion words:", len(CONCLUSION.split()))
