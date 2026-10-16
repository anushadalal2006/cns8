"""Build the two Word documents (and PDFs via LibreOffice) from the measured data.

  report/Exp08_Tables_and_Graphs_<UID>.docx   Tables 1-7 and Graphs 1-7 (editable tables)
  report/Exp08_Screenshots_<UID>.docx         terminal screenshots in order of execution

usage: python3 tools/make_docs.py   (run tools/render_screenshots.py first)
"""
import csv
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from render_screenshots import SHOTS  # noqa: E402

NAME, UID, DATE = "Anusha Dalal", "2024300032", "16 October 2026"
TITLE = ("Experiment 8: Evaluation of Classical and Post-Quantum Digital Signature Algorithms "
         "using Docker-Based Client-Server Communication")
ALGS = ["ecdsa", "rsapss", "mldsa", "slhdsa"]
LABEL = {"ecdsa": "ECDSA", "rsapss": "RSA-PSS", "mldsa": "ML-DSA", "slhdsa": "SLH-DSA"}

# ------------------------------------------------------------------ data
with open("sender/results.csv") as f:
    SENDER = list(csv.DictReader(f))
T = {(r["algorithm"], r["operation"]): (float(r["mean_ms"]), float(r["std_ms"])) for r in SENDER}


def ms(a, op):
    return T[(a, op)][0]


sizes = {}
for line in open("logs/G_sizes.log"):
    m = re.match(r"(\w+) public=(\d+) private=(\d+) sig=(\d+) doc=(\d+)", line)
    if m:
        a, pk, sk, sg, doc = m.groups()
        sizes[a] = dict(pk=int(pk), sk=int(sk), sig=int(sg), doc=int(doc))
for a in ALGS:
    sizes[a]["tx"] = sizes[a]["doc"] + sizes[a]["sig"] + sizes[a]["pk"]

results3, cur = {}, None
for line in open("logs/H_tamper.log"):
    m = re.search(r"verify\.py (\w+) .*# Test (\d)", line)
    if m:
        cur = (m.group(2), m.group(1))
    elif "RESULT" in line and cur:
        results3[cur] = "PASS" if "PASS" in line else "FAIL"
        cur = None

DOC_TEXT = open("sender/outbox/document.txt").read().splitlines()
TAMPERED = open("receiver/inbox/document.txt.tampered", encoding="utf-8").read().splitlines()
ENV = "; ".join(l.strip() for l in open("sender/environment.txt") if l.strip())


# ------------------------------------------------------------------ docx helpers
def new_doc():
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    return d


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def para(d, text, bold=False, italic=False, size=None, align=None, after=None):
    p = d.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if align is not None:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def header(d, subtitle):
    para(d, "Sardar Patel Institute of Technology - Department of Computer Engineering",
         size=10, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(d, "Cryptography and Network Security Lab", size=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, TITLE, bold=True, size=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, subtitle, bold=True, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(d, f"Name: {NAME}        UID: {UID}        Date: {DATE}", size=11,
         align=WD_ALIGN_PARAGRAPH.CENTER)


def table(d, title, head, rows, widths=None, size=9.5):
    if title:
        d.add_heading(title, level=2)
    t = d.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        c = t.rows[0].cells[i]
        run = c.paragraphs[0].add_run(str(h)); run.bold = True; run.font.size = Pt(size)
        shade(c, "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            run = cells[i].paragraphs[0].add_run(str(v)); run.font.size = Pt(size)
            run.bold = i == 0
    if widths:
        t.autofit = False
        for i, w in enumerate(widths):
            t._tbl.tblGrid.gridCol_lst[i].w = Cm(w)
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t


def note(d, text):
    para(d, text, italic=True, size=9)


def picture(d, path, max_w=17.0, max_h=23.0):
    w, h = Image.open(path).size
    width = min(max_w, max_h * w / h)
    d.add_picture(path, width=Cm(width))
    d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def page_break(d):
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def save(d, path):
    d.save(path)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir",
                    os.path.dirname(path), path], check=True, capture_output=True)
    print("wrote", path, "and", path.replace(".docx", ".pdf"))


# ================================================================== 1. Tables and graphs
d = new_doc()
header(d, "Tables and Graphs")

table(d, "Table 1: Execution environment",
      ["CPU model", "Cores / threads", "RAM", "Host OS", "Docker version", "cryptography version",
       "liboqs version"],
      [["Intel(R) Xeon(R) Processor @ 2.80GHz", "4 / 4", "15 GiB", "Ubuntu 24.04.5 LTS (kernel 6.18.44)",
        "29.8.2", "50.0.2", "0.16.0 (liboqs-python 0.16.0.1)"]], size=9)
note(d, "Containers: image sig-lab:1.0 (Ubuntu 24.04.5 LTS, Python 3.12.3); sender 172.18.0.2, receiver 172.18.0.3 on sig-net.")

table(d, "Table 2: Key, signature and transmission sizes (bytes)",
      ["Algorithm", "Public key", "Private key", "Signature", "Document", "Transmitted (doc + sig + public key)"],
      [[LABEL[a], sizes[a]["pk"], sizes[a]["sk"], sizes[a]["sig"], sizes[a]["doc"], sizes[a]["tx"]] for a in ALGS])
note(d, "Sizes from stat -c %s. ML-DSA-44 (1312 / 2560 / 2420) and SLH-DSA-SHA2-128s (32 / 64 / 7856) match the "
        "standards' raw sizes; ECDSA and RSA-PSS keys are DER encoded (PKCS#8 / SubjectPublicKeyInfo).")

names = {"1": "1. Original document + original signature", "2": "2. One character of document changed",
         "3": "3. One bit of signature changed", "4": "4. Correct document + signature, wrong public key"}
table(d, "Table 3: Verification results",
      ["Test case", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA", "Expected"],
      [[nm] + [results3[(t, a)] for a in ALGS] + ["PASS" if t == "1" else "FAIL"] for t, nm in names.items()],
      widths=[6.6, 2, 2, 2, 2, 2])
para(d, "")
table(d, None, ["Document", "Text (tamper index = UID mod 59 = 2: 'a' (0x61) changed to '`' (0x60))"],
      [["Original", f"{DOC_TEXT[0]} / {DOC_TEXT[1]}"], ["Modified", f"{TAMPERED[0]} / {TAMPERED[1]}"]],
      widths=[2.6, 14])
note(d, "Test 3 flipped bit 0 of byte UID mod length: ECDSA byte 4, RSA-PSS byte 256, ML-DSA byte 1492, "
        "SLH-DSA byte 5232. Test 4, ECDSA = PASS: the ECDSA private key is derived from the UID (Section 5) and "
        "both containers have the same UID, so the receiver's key pair is identical to the sender's (confirmed "
        "with cmp) - the 'wrong' key is actually the right one.")

table(d, "Table 4: Timing results (milliseconds, mean ± standard deviation)",
      ["Algorithm", "Key generation", "Signing", "Verification", "Sign / Verify ratio"],
      [[LABEL[a]] + [f"{T[(a, op)][0]:.4f} ± {T[(a, op)][1]:.4f}" for op in ("keygen", "sign", "verify")]
       + [f"{ms(a, 'sign') / ms(a, 'verify'):.2f}"] for a in ALGS])
note(d, "Sender, median of three bench.py runs (run 2); 1024-byte UID-seeded message; 5 warm-up calls, then "
        "key generation x20, signing x100, verification x100, each call timed with time.perf_counter_ns().")
page_break(d)

d.add_heading("Graphs (Section 15)", level=1)
graphs = [
    ("graph1_key_sizes.png", "Graph 1: Public-key and private-key size (bytes, log scale)"),
    ("graph2_signature_size.png", "Graph 2: Signature size (bytes, log scale)"),
    ("graph3_signing_time.png", "Graph 3: Average signing time (ms, log scale); error bar = standard deviation, n = 100"),
    ("graph4_verification_time.png", "Graph 4: Average verification time (ms, log scale); error bar = standard deviation, n = 100"),
    ("graph5_keygen_time.png", "Graph 5: Average key-generation time (ms, log scale), n = 20"),
    ("graph6_transmitted_bytes.png", "Graph 6: Total transmitted bytes = document + signature + public key"),
    ("graph7_relative_to_ecdsa.png", "Graph 7 (extra): Sign, verify and key-generation time relative to ECDSA (ECDSA = 1, log scale)"),
]
for g, cap in graphs:
    picture(d, os.path.join("graphs", g), max_w=15.5, max_h=10.5)
    para(d, f"{cap}. Environment: {ENV}. UID {UID}.", italic=True, size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
page_break(d)


def rel(vals):
    return ["1.0"] + [f"{vals[a] / vals['ecdsa']:.2f}" for a in ALGS[1:]]


table(d, "Table 5: Relative comparison with ECDSA as the baseline (measured value / ECDSA value)",
      ["Parameter", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA"],
      [["Public key size"] + rel({a: sizes[a]["pk"] for a in ALGS}),
       ["Signature size"] + rel({a: sizes[a]["sig"] for a in ALGS}),
       ["Signing time"] + rel({a: ms(a, "sign") for a in ALGS}),
       ["Verification time"] + rel({a: ms(a, "verify") for a in ALGS}),
       ["Key generation time"] + rel({a: ms(a, "keygen") for a in ALGS}),
       ["Total transmitted bytes"] + rel({a: sizes[a]["tx"] for a in ALGS})])

f3 = lambda a, op: f"{ms(a, op):.3f} ms"  # noqa: E731
table(d, "Table 6: Property comparison",
      ["Property", "ECDSA", "RSA-PSS", "ML-DSA", "SLH-DSA"],
      [["Standard", "FIPS 186-5", "RFC 8017, FIPS 186-5", "FIPS 204 (2024)", "FIPS 205 (2024)"],
       ["Security parameter / category", "P-256 (secp256r1), ~128-bit", "RSA-3072, ~128-bit",
        "ML-DSA-44, NIST category 2", "SLH-DSA-SHA2-128s, NIST category 1"],
       ["Mathematical basis", "Elliptic-curve discrete logarithm", "Integer factorisation",
        "Module lattices (Module-LWE / SIS)", "Hash functions only (hypertree of Merkle / FORS trees)"],
       ["Resistant to Shor's algorithm?", "No", "No", "Yes", "Yes"],
       ["Public key size (small / medium / large)", f"Small ({sizes['ecdsa']['pk']} B)",
        f"Medium ({sizes['rsapss']['pk']} B)", f"Large ({sizes['mldsa']['pk']} B)", f"Small ({sizes['slhdsa']['pk']} B)"],
       ["Signature size", f"Small ({sizes['ecdsa']['sig']} B)", f"Medium ({sizes['rsapss']['sig']} B)",
        f"Large ({sizes['mldsa']['sig']} B)", f"Very large ({sizes['slhdsa']['sig']} B)"],
       ["Fastest operation", f"Signing ({f3('ecdsa', 'sign')})", f"Verification ({f3('rsapss', 'verify')})",
        f"Verification ({f3('mldsa', 'verify')})", f"Verification ({f3('slhdsa', 'verify')})"],
       ["Slowest operation", f"Verification ({f3('ecdsa', 'verify')})", f"Key generation ({ms('rsapss', 'keygen'):.1f} ms)",
        f"Signing ({f3('mldsa', 'sign')})", f"Signing ({ms('slhdsa', 'sign'):.1f} ms)"],
       ["Maturity / deployment", "Very mature; TLS, SSH, code signing, smart cards",
        "Very mature (decades); PKI, TLS, document signing", "New (2024); being rolled out, often in hybrid mode",
        "New (2024), conservative SPHINCS+ design; firmware, root keys"],
       ["Main advantage", "Smallest keys and signatures, fast signing", "Very fast verification, universal support",
        "Quantum-safe and fast (fastest verification measured)", "Quantum-safe, relies only on hash functions; 32 B public key"],
       ["Main disadvantage", "Broken by Shor; needs secure nonces", "Broken by Shor; slow key generation, larger keys",
        "Large public key and signature", "Very slow signing, very large signature"]],
      widths=[3.4, 3.3, 3.3, 3.3, 3.4], size=8.5)

table(d, "Table 7: Which algorithm would you choose?",
      ["Scenario", "Chosen algorithm", "Justification"],
      [["Signing a firmware image for a device that stays in the field for 15 years", "SLH-DSA",
        "Must outlive quantum computers; hash-only security is the most conservative, signing is rare and verification is cheap."],
       ["Authenticating a short sensor message on a very low-bandwidth link", "ECDSA",
        f"Smallest overhead ({sizes['ecdsa']['sig']} B signature, {sizes['ecdsa']['pk']} B key) and the data is short-lived."],
       ["Signing a TLS server certificate in 2026 on a general web server", "ECDSA (P-256)",
        "Universally supported, tiny handshake; certificates live < 1 year, with a planned move to hybrid ECDSA + ML-DSA."],
       ["Root key of a certificate authority that must remain trusted for 30 years", "SLH-DSA",
        "Rarely signs, so slow signing and large signatures are acceptable; security rests only on hash functions."],
       ["High-volume server that verifies millions of signatures per second", "ML-DSA",
        f"Fastest verification measured ({f3('mldsa', 'verify')}) and quantum-safe."]],
      widths=[5.2, 2.8, 9], size=9)

os.makedirs("report", exist_ok=True)
save(d, f"report/Exp08_Tables_and_Graphs_{UID}.docx")

# ================================================================== 2. Screenshots
d = new_doc()
header(d, "Screenshots (in order of execution)")
for i, (name, heading, *_rest) in enumerate(SHOTS, 1):
    h = d.add_heading(f"Screenshot {i}: {heading}", level=3)
    h.paragraph_format.page_break_before = i > 1
    picture(d, os.path.join("screenshots", name), max_w=17.0, max_h=22.5)
save(d, f"report/Exp08_Screenshots_{UID}.docx")
