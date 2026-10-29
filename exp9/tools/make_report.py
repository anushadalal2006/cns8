"""Build the Lab 9 report (.docx + .pdf) and the submission zip from the student's run.

usage: python3 exp9/tools/make_report.py   (after render_terminal.py and wireshark_views.py)
"""
import os
import re
import shutil
import subprocess
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
EXP9 = os.path.dirname(HERE)
TERM = os.path.join(EXP9, "terminal_screenshots")
WS = os.path.join(EXP9, "wireshark_screenshots")
SUB = os.path.join(EXP9, "submission")
REPORT = os.path.join(EXP9, "report")
NAME, UID, DATE = "Anusha Dalal", "2024300032", "29 October 2026"
BASE = f"Lab9_X509_PKI_{UID}_{NAME.replace(' ', '_')}"

manifest = {}
for line in open(os.path.join(TERM, "manifest.txt")):
    heading, files = line.rstrip("\n").split("\t")
    manifest[files.split(",")[0].rsplit("_", 1)[0] if re.search(r"_\d\.png$", files.split(",")[0])
             else files.split(",")[0][:-4]] = (heading, files.split(","))

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(11)
for s in doc.sections:
    s.left_margin = s.right_margin = Cm(2)
    s.top_margin = s.bottom_margin = Cm(1.8)
fig_no = 0


def para(text, bold=False, italic=False, size=None, align=None, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold, run.italic = bold, italic
    if size:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)
    if align is not None:
        p.alignment = align
    return p


def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def picture(path, max_w=17.0, max_h=21.5):
    w, h = Image.open(path).size
    doc.add_picture(path, width=Cm(min(max_w, max_h * w / h)))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def figure(path, caption, max_w=17.0, max_h=21.5):
    global fig_no
    fig_no += 1
    h = doc.add_paragraph()
    r = h.add_run(f"Figure {fig_no}: {caption}")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    h.paragraph_format.keep_with_next = True
    picture(path, max_w, max_h)


def term(stem):
    heading, files = manifest[stem]
    for k, f in enumerate(files):
        cap = heading if len(files) == 1 else f"{heading} ({k + 1}/{len(files)})"
        figure(os.path.join(TERM, f), cap)


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def table(head, rows, widths, size=9):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(head):
        run = t.rows[0].cells[i].paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(size)
        shade(t.rows[0].cells[i], "D9E2F3")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            run = cells[i].paragraphs[0].add_run(v); run.font.size = Pt(size)
    t.autofit = False
    for i, w in enumerate(widths):
        t._tbl.tblGrid.gridCol_lst[i].w = Cm(w)
        for r in t.rows:
            r.cells[i].width = Cm(w)
    for r in t.rows[:-1]:
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True
    doc.add_paragraph()


def answers_box(title):
    t = doc.add_table(rows=1, cols=1)
    t.style = "Table Grid"
    c = t.rows[0].cells[0]
    run = c.paragraphs[0].add_run(f"\n{title}\n[handwritten answers attached]\n")
    run.italic = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(c, "F3F3F3")
    doc.add_paragraph()



ANSWERS = open(os.path.join(EXP9, "ANSWERS_TO_HANDWRITE_LAB9.md"), encoding="utf-8").read()


def rich(p, text, size=10.5):
    """add **bold**, *italic* and `code` runs to paragraph p"""
    for tok in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)", text):
        if not tok:
            continue
        if tok.startswith("**"):                      # bold, possibly containing `code`
            for k, part_ in enumerate(re.split(r"`([^`]+)`", tok[2:-2])):
                if not part_:
                    continue
                r = p.add_run(part_); r.bold = True; r.font.size = Pt(size)
                if k % 2:
                    r.font.name = "Consolas"
                    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            continue
        elif tok.startswith("`"):
            r = p.add_run(tok[1:-1]); r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif tok.startswith("*"):
            r = p.add_run(tok[1:-1]); r.italic = True
        else:
            r = p.add_run(tok)
        r.font.size = Pt(size)


def answers(task):
    """render the '## Task X ...' section of the answers file into the report"""
    m = re.search(r"^## (" + re.escape(task) + r"[^\n]*)\n(.*?)(?=^---$|^## )", ANSWERS, re.S | re.M)
    doc.add_heading(f"Answers: {m.group(1)}", level=2)
    lines = m.group(2).strip("\n").split("\n")
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):                       # markdown table
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"-+", c) for c in cells):
                    tbl.append(cells)
                i += 1
            n = len(tbl[0])
            t = doc.add_table(rows=0, cols=n)
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r_i, cells in enumerate(tbl):
                row = t.add_row().cells
                for c_i, val in enumerate(cells):
                    rich(row[c_i].paragraphs[0], val, 9)
                    if r_i == 0:
                        for run in row[c_i].paragraphs[0].runs:
                            run.bold = True
                        shade(row[c_i], "D9E2F3")
            widths = [3.6] + [(17 - 3.6) / (n - 1)] * (n - 1)
            t.autofit = False
            for c_i, w in enumerate(widths):
                t._tbl.tblGrid.gridCol_lst[c_i].w = Cm(w)
                for r in t.rows:
                    r.cells[c_i].width = Cm(w)
            doc.add_paragraph()
            continue
        mb = re.match(r"^(\s*)- (.*)$", line)
        mn = re.match(r"^(\d+)\. (.*)$", line)
        if mb:
            p = doc.add_paragraph(style="List Bullet 2" if mb.group(1) else "List Bullet")
            rich(p, mb.group(2))
        elif mn:
            p = doc.add_paragraph()
            rich(p, mn.group(1) + ". " + mn.group(2))
            p.paragraph_format.left_indent = Cm(0.6)
        else:
            p = doc.add_paragraph()
            rich(p, line)
        p.paragraph_format.space_after = Pt(3)
        i += 1


def part(title, intro):
    page_break()
    doc.add_heading(title, level=1)
    para(intro, size=10.5)


# ------------------------------------------------------------------ title
para("Bharatiya Vidya Bhavan's", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=10)
para("SARDAR PATEL INSTITUTE OF TECHNOLOGY", bold=True, size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
para("Department of Computer Engineering — Cryptography and Network Security Lab",
     align=WD_ALIGN_PARAGRAPH.CENTER, size=10.5)
para("Lab 9: Building and Securing a Two-Tier PKI with X.509", bold=True, size=16,
     align=WD_ALIGN_PARAGRAPH.CENTER)
table(["Field", "Value"], [["Name", NAME], ["UID", UID], ["Date of run", DATE],
                           ["Environment", "Ubuntu on WSL2 (Windows), nginx/1.24.0 (Ubuntu), curl/8.5.0, "
                                           "OpenSSL and tshark from the Ubuntu packages"]], [4, 13], size=10)
doc.add_heading("Aim", level=2)
para("To build a two-tier certificate authority (Root CA and Intermediate CA) with OpenSSL, issue server and "
     "client certificates, configure nginx for mutual TLS, revoke a certificate and enforce it with a CRL and an "
     "OCSP responder, analyse TLS 1.2 and TLS 1.3 handshakes, and reproduce the common certificate-validation "
     "failures.")
doc.add_heading("Report layout", level=2)
table(["Section", "Contents"],
      [["Setup", "Folder layout and CA databases"],
       ["Part A", "Root CA: configuration, key, self-signed certificate, inspection"],
       ["Part B", "Intermediate CA: CSR signed by the Root, chain file, pathlen and leaf-signing tests"],
       ["Part C", "Server and client certificates, verification, index.txt, purpose and SAN"],
       ["Part D", "Mutual TLS with nginx: valid client, missing client certificate, s_client"],
       ["Part E", "Revocation: CRLs, nginx ssl_crl, OCSP responder"],
       ["Part F", "TLS 1.2 and TLS 1.3 handshake captures (tshark output and Wireshark views)"],
       ["Part G", "Six failure scenarios with curl, OpenSSL and nginx error.log evidence"],
       ["Reflection", "Handwritten (150–250 words)"],
       ["Submission", "Files included in the zip (no private keys)"]], [3, 14], size=10)
para("The answers to the Task questions follow each Part. The reflection is handwritten and attached at the "
     "marked place.", italic=True, size=10)

part("Setup", "Working directory $PKI = ~/pki-lab with separate folders for the Root and the Intermediate CA. "
     "index.txt is each CA's database, serial the next serial number and crlnumber the next CRL number.")
term("01_setup")

# ------------------------------------------------------------------ Part A
part("Part A — Root CA", "root/openssl.cnf defines the Root CA (files, policy, default hash sha256) and the "
     "extension profiles v3_root_ca and v3_intermediate_ca. The Root key is 4096-bit RSA and the certificate is "
     "self-signed for 10 years (3650 days).")
term("A1_root_config")
term("A2_root_key_cert")
term("A3_root_inspect")
answers("Task A")

# ------------------------------------------------------------------ Part B
part("Part B — Intermediate CA", "The Intermediate creates its own 4096-bit key and a CSR, and the Root signs it "
     "with the v3_intermediate_ca profile (CA:TRUE, pathlen:0) for 5 years. intermediate/openssl.cnf also holds the "
     "server_cert and client_cert profiles, the SAN list and the OCSP URL.")
term("B1_intermediate_config")
term("B2_intermediate_sign")
term("B3_chain_verify")
term("B4_intermediate_inspect")
term("B5_task_demos")
answers("Task B")

# ------------------------------------------------------------------ Part C
part("Part C — Server and client certificates", "The Intermediate signs a 2048-bit server certificate "
     "(server_cert: serverAuth, SAN server.lab.local / localhost / 127.0.0.1, OCSP URL) and client certificates "
     "for Alice and Bob (client_cert: clientAuth). server.fullchain.pem = server certificate + Intermediate.")
term("C1_server_cert")
term("C2_client_certs")
term("C3_verify")
term("C4_task")
term("C5_server_inspect")
answers("Task C")

# ------------------------------------------------------------------ Part D
part("Part D — Mutual TLS with nginx", "nginx listens on 8443 with the server full chain, requires a client "
     "certificate (ssl_verify_client on) and verifies it against the CA chain (ssl_client_certificate, "
     "ssl_verify_depth 2). The response prints the client DN, the verification result and the TLS version.")
term("D1_nginx_config")
term("D2_test")
term("D3_s_client")
answers("Task D")

# ------------------------------------------------------------------ Part E
part("Part E — Revocation (CRL and OCSP)", "Bob's certificate (serial 1002) is revoked with reason "
     "keyCompromise, both CAs publish a CRL, and nginx enforces the combined CRL. An OCSP responder answers "
     "live status queries from the Intermediate's index.txt.")
term("E1_revoke_crl")
term("E1b_verify_crl")
term("E2_nginx_crl")
term("E2b_without_root_crl")
term("E3_responder")
term("E3_alice_top")
term("E3_alice_end")
term("E3_bob_top")
term("E3_bob_end")
term("E3_aia")
answers("Task E")

# ------------------------------------------------------------------ Part F
part("Part F — TLS handshake analysis", "Both handshakes were captured on the loopback interface (lo) with the "
     "capture filter tcp port 8443, using Alice's certificate. For TLS 1.3 the session secrets were written to "
     "keys.log (SSLKEYLOGFILE) during the same captured connection, so that the capture can be decrypted.")
doc.add_heading("F.1 TLS 1.2", level=2)
term("F1_capture12")
term("F1_tls12")
ref = os.environ.get("WS_REF_DIR", "")
if ref and os.path.exists(os.path.join(ref, "4.png")):
    figure(os.path.join(ref, "4.png"), "Wireshark: tls12.pcapng with display filter tls.handshake", max_h=6)
    figure(os.path.join(ref, "5.png"), "Wireshark: frame 4, Client Hello (TLS 1.2, 28 cipher suites)", max_h=20)
ws = [("F1_02_serverhello_tls12.png", "Wireshark view: frame 6, Server Hello - selected cipher suite"),
      ("F1_03_certificate_tls12.png", "Wireshark view: frame 6, server Certificate - issuer, validity, subject, "
                                      "extensions (EKU, SAN, OCSP URL)"),
      ("F1_04a_certificate_request_tls12.png", "Wireshark view: frame 6, Certificate Request with the acceptable CA "
                                               "names"),
      ("F1_04b_certificate_verify_tls12.png", "Wireshark view: frame 8, client Certificate (CN=alice) and "
                                              "Certificate Verify")]
for f, cap in ws:
    figure(os.path.join(WS, f), cap, max_h=21.5)
doc.add_heading("F.2 TLS 1.3 without the key log", level=2)
term("F2_capture13")
term("F2_tls13")
figure(os.path.join(WS, "F2_05_tls13_encrypted.png"), "Wireshark view: tls13.pcapng, filter tls - after the "
       "Server Hello everything is Application Data", max_h=8)
figure(os.path.join(WS, "F2_06_tls13_no_certificate.png"), "Wireshark view: tls13.pcapng, filter "
       "tls.handshake.type == 11 - no Certificate message is visible", max_h=6)
doc.add_heading("F.3 TLS 1.3 decrypted with keys.log", level=2)
term("F3_tls13_decrypted")
figure(os.path.join(WS, "F3_07_tls13_decrypted.png"), "Wireshark view: tls13.pcapng with keys.log, filter "
       "tls.handshake - Certificate, Certificate Verify and Finished decrypted", max_h=7)
figure(os.path.join(WS, "F3_08_tls13_certificate.png"), "Wireshark view: frame 8 with keys.log - client "
       "Certificate (CN=alice, clientAuth) and Certificate Verify", max_h=21.5)
answers("Task F")

# ------------------------------------------------------------------ Part G
part("Part G — Failure scenarios", "Each failure was created on purpose and tested with curl, OpenSSL and the "
     "nginx error.log (level info). nginx had ssl_crl enabled from Part E during all scenarios.")
for stem in ("G1_expired", "G2_untrusted", "G3_wrong_purpose", "G4_hostname", "G5_missing_intermediate",
             "G6_revoked", "G7_errorlog"):
    term(stem)
doc.add_heading("Observed results", level=2)
table(["#", "Scenario", "curl", "nginx error.log", "OpenSSL"],
      [["1", "Expired client (eve)", "HTTP 400 The SSL certificate error", "10: certificate has expired",
        "verify: error 10"],
       ["2", "Untrusted self-signed client (mallory)", "HTTP 400 The SSL certificate error",
        "3: unable to get certificate CRL (ssl_crl on)", "verify: error 18 self-signed certificate"],
       ["3", "Wrong purpose (server cert as client)", "HTTP 400 The SSL certificate error",
        "26: unsuitable certificate purpose", "verify -purpose sslclient: error 26"],
       ["4", "Hostname mismatch (wrong.lab.local)", "curl (60): no alternative certificate subject name matches",
        "client closed connection (no verify error)", "s_client -verify_hostname: 62 hostname mismatch"],
       ["5", "Missing intermediate (leaf only)", "chain in --cacert: works; Root only: (60) unable to get local "
        "issuer certificate", "SSL alert 48: unknown ca", "s_client: 20, then 21"],
       ["6", "Revoked (bob)", "HTTP 400 The SSL certificate error", "23: certificate revoked",
        "verify -crl_check: error 23"]],
      [0.7, 3.6, 4.4, 4.3, 4.0], size=8.5)
answers("Task G")

# ------------------------------------------------------------------ reflection and submission
part("Reflection", "What is the single most common reason real-world TLS deployments fail certificate validation, "
     "and how would you prevent it? (150 to 250 words)")
answers_box("Reflection")

page_break()
doc.add_heading("Submission files", level=1)
para("The zip file contains this report and the following files from ~/pki-lab. No private keys (*.key.pem) are "
     "included.", size=10.5)
files = sorted(os.path.relpath(os.path.join(dp, f), SUB) for dp, _, fs in os.walk(SUB) for f in fs)
purpose = {"openssl.cnf": "CA configuration", "index.txt": "CA database", "nginx.conf": "nginx configuration",
           "error.log": "nginx error log (level info)", "error_log_excerpts.txt": "client-certificate errors only",
           "access.log": "nginx access log", "keys.log": "TLS 1.3 session secrets (SSLKEYLOGFILE)",
           "tls12.pcapng": "TLS 1.2 capture", "tls13.pcapng": "TLS 1.3 capture"}
rows = []
for f in files:
    base = os.path.basename(f)
    desc = purpose.get(base) or ("CRL" if base.endswith(".crl.pem") else "certificate chain" if "chain" in base
                                 else "certificate")
    rows.append([f, desc])
table(["File", "Contents"], rows, [9.5, 7.5], size=9)

os.makedirs(REPORT, exist_ok=True)
docx_path = os.path.join(REPORT, BASE + ".docx")
doc.save(docx_path)
subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", REPORT, docx_path],
               check=True, capture_output=True)
pdf_path = docx_path[:-5] + ".pdf"
print("wrote", docx_path, "and", pdf_path)

# ------------------------------------------------------------------ zip (report + files, no private keys)
zip_path = os.path.join(REPORT, BASE + ".zip")
with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
    z.write(pdf_path, os.path.basename(pdf_path))
    for f in files:
        if f.endswith(".key.pem"):
            sys.exit("private key found - aborting")
        z.write(os.path.join(SUB, f), os.path.join("pki-lab", f))
print("wrote", zip_path)
